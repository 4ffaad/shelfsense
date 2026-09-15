"""Generate the illustrated ShelfSense learning plan."""

# The document text deliberately uses long prose lines; Word handles wrapping.
# ruff: noqa: E501

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs"
ASSETS = OUT / "learning-assets"
ASSETS.mkdir(parents=True, exist_ok=True)

NAVY = (15, 23, 42)
BLUE = (37, 99, 235)
TEAL = (13, 148, 136)
PURPLE = (124, 58, 237)
ORANGE = (234, 88, 12)
GRAY = (71, 85, 105)
LIGHT = (241, 245, 249)
WHITE = (255, 255, 255)


def font(size, bold=False):
    candidates = [
        "/System/Library/Fonts/Supplemental/Arial Bold.ttf" if bold else "/System/Library/Fonts/Supplemental/Arial.ttf",
        "/System/Library/Fonts/SFNS.ttf",
    ]
    for candidate in candidates:
        if Path(candidate).exists():
            return ImageFont.truetype(candidate, size)
    return ImageFont.load_default()


def centered(draw, box, text, f, fill=WHITE):
    x1, y1, x2, y2 = box
    bb = draw.multiline_textbbox((0, 0), text, font=f, align="center", spacing=3)
    draw.multiline_text(((x1 + x2 - (bb[2]-bb[0]))/2, (y1 + y2 - (bb[3]-bb[1]))/2), text, font=f, fill=fill, align="center", spacing=3)


def pipeline_png():
    w, h = 1500, 470
    im = Image.new("RGB", (w, h), NAVY)
    d = ImageDraw.Draw(im)
    title = font(30, True)
    label = font(20, True)
    small = font(15)
    d.text((40, 25), "ShelfSense: the product pipeline", font=title, fill=WHITE)
    steps = [
        ("1\nImage", BLUE, "input"),
        ("2\nValidate", TEAL, "safe bytes"),
        ("3\nDetect", PURPLE, "where?"),
        ("4\nIdentify", ORANGE, "which SKU?"),
        ("5\nCount", BLUE, "how many?"),
        ("6\nEvents", TEAL, "what action?"),
    ]
    x0, y, bw, bh, gap = 50, 150, 190, 155, 48
    for i, (txt, color, sub) in enumerate(steps):
        x = x0 + i * (bw + gap)
        d.rounded_rectangle((x, y, x+bw, y+bh), radius=18, fill=color)
        centered(d, (x, y+15, x+bw, y+105), txt, label)
        centered(d, (x, y+105, x+bw, y+145), sub, small, fill=(226,232,240))
        if i < len(steps)-1:
            ax = x+bw+8
            ay = y+bh//2
            d.line((ax, ay, ax+gap-16, ay), fill=(148,163,184), width=5)
            d.polygon([(ax+gap-16, ay), (ax+gap-30, ay-10), (ax+gap-30, ay+10)], fill=(148,163,184))
    d.text((52, 380), "We implement and test these stages one at a time—not all at once.", font=font(20), fill=(203,213,225))
    im.save(ASSETS / "pipeline.png")


def roadmap_png():
    w, h = 1500, 600
    im = Image.new("RGB", (w, h), NAVY)
    d = ImageDraw.Draw(im)
    d.text((40, 25), "Learning roadmap", font=font(30, True), fill=WHITE)
    phases = [
        ("0", "Foundation", "Python • API • tests", BLUE),
        ("1", "Images", "uploads • validation", TEAL),
        ("2", "Detection", "boxes • confidence", PURPLE),
        ("3", "Identity", "SKU • unknown", ORANGE),
        ("4", "Inventory", "counts • alerts", BLUE),
        ("5", "Dashboard", "UI • product", TEAL),
    ]
    x0, y, bw, bh, gap = 42, 145, 210, 245, 28
    for i, (num, name, sub, color) in enumerate(phases):
        x = x0 + i*(bw+gap)
        d.rounded_rectangle((x, y, x+bw, y+bh), radius=16, fill=(30,41,59), outline=color, width=4)
        d.ellipse((x+72, y+25, x+138, y+91), fill=color)
        centered(d, (x+72, y+25, x+138, y+91), num, font(28, True))
        centered(d, (x+10, y+110, x+bw-10, y+165), name, font(22, True), fill=WHITE)
        centered(d, (x+10, y+170, x+bw-10, y+225), sub, font(16), fill=(203,213,225))
        if i < len(phases)-1:
            ax, ay = x+bw+4, y+bh//2
            d.line((ax, ay, ax+gap-8, ay), fill=(148,163,184), width=4)
            d.polygon([(ax+gap-8, ay), (ax+gap-20, ay-9), (ax+gap-20, ay+9)], fill=(148,163,184))
    d.rounded_rectangle((60, 455, 1440, 535), radius=15, fill=(30,41,59))
    d.text((90, 480), "Rule: finish one measurable vertical slice before adding the next layer.", font=font(22, True), fill=(226,232,240))
    im.save(ASSETS / "roadmap.png")


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def set_cell_text(cell, text, bold=False, color="1E293B"):
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(text)
    r.bold = bold
    r.font.color.rgb = RGBColor.from_string(color)
    r.font.size = Pt(9.5)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def make_doc():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.65)
    sec.bottom_margin = Inches(0.65)
    sec.left_margin = Inches(0.8)
    sec.right_margin = Inches(0.8)

    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10.5)
    for name, size, color in [("Title", 30, "0F172A"), ("Heading 1", 20, "1D4ED8"), ("Heading 2", 14, "0F766E")]:
        styles[name].font.name = "Aptos Display"
        styles[name].font.size = Pt(size)
        styles[name].font.color.rgb = RGBColor.from_string(color)

    p = doc.add_paragraph()
    p.style = "Title"
    p.add_run("ShelfSense\n")
    r = p.add_run("A visual learning plan for building a retail inventory system")
    r.font.size = Pt(15)
    r.font.color.rgb = RGBColor.from_string("475569")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Learn computer vision by building the system one understandable piece at a time.")
    r.italic = True
    r.font.size = Pt(12)

    doc.add_picture(str(ASSETS / "pipeline.png"), width=Inches(6.65))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_heading("1. What are we building?", level=1)
    doc.add_paragraph("ShelfSense is a local-first computer-vision inventory tool. It receives a shelf image, finds visible products, identifies known products, counts them, and reports inventory conditions such as low stock or out of stock.")
    doc.add_paragraph("The important learning decision is to avoid building the whole platform immediately. We will build a sequence of small behaviors that can be run, tested, measured, and explained.")

    doc.add_heading("2. The plan in one sentence", level=1)
    p = doc.add_paragraph()
    r = p.add_run("Turn one shelf image into one trustworthy inventory observation.")
    r.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = RGBColor.from_string("0F766E")

    doc.add_heading("3. The roadmap", level=1)
    doc.add_picture(str(ASSETS / "roadmap.png"), width=Inches(6.65))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for c, text in zip(table.rows[0].cells, ["Phase", "What we build", "What you learn"], strict=True):
        set_cell_text(c, text, True, "FFFFFF")
        shade(c, "1D4ED8")
    rows = [
        ("0 — Foundation", "FastAPI app, settings, health route, tests", "Project structure, HTTP, testing, tooling"),
        ("1 — Image ingestion", "Upload and validate an image", "Files, pixels, validation, errors"),
        ("2 — Detection", "Find products with bounding boxes", "Models, confidence, IoU, precision/recall"),
        ("3 — Identification", "Map product crops to 3–5 SKUs", "Classification, datasets, unknown results"),
        ("4 — Inventory", "Count SKUs and create events", "Domain models, databases, business rules"),
        ("5 — Dashboard", "Show observations and alerts", "Frontend, API contracts, product design"),
    ]
    for row in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            set_cell_text(cells[i], text)
            if len(table.rows) % 2 == 0:
                shade(cells[i], "F8FAFC")

    doc.add_heading("4. How each coding session works", level=1)
    steps = [
        "Choose one narrow behavior.",
        "Write down its input and expected output.",
        "Implement the smallest version.",
        "Write tests for normal and invalid input.",
        "Run it on a real example.",
        "Measure the result and inspect failures.",
        "Explain what happened before changing it.",
    ]
    for s in steps:
        doc.add_paragraph(s, style="List Number")
    doc.add_paragraph("This is the core skill: not merely writing code, but connecting a requirement to an implementation, a test, and evidence.")

    doc.add_heading("5. The first milestone", level=1)
    doc.add_paragraph("Our first milestone is intentionally small:")
    for s in [
        "One uploaded shelf image",
        "One configured shelf",
        "Three to five known products",
        "Local inference on your M3",
        "SQLite persistence",
        "A FastAPI endpoint returning structured results",
        "Tests and a small evaluation report",
    ]:
        doc.add_paragraph(s, style="List Bullet")
    doc.add_paragraph("We will not start with live video, a multi-store system, cloud deployment, or a polished dashboard. Those are later layers, not prerequisites for learning the fundamentals.")

    doc.add_heading("6. What the technical pieces mean", level=1)
    concepts = [
        ("Detection", "Where are the products? The result is a set of bounding boxes."),
        ("Identification", "What product is inside each detected box? The result is a SKU or unknown."),
        ("Aggregation", "How many detections belong to each SKU?"),
        ("Observation", "What did this particular image show at this time?"),
        ("Event", "What should the business know or act on, such as low stock?"),
        ("Abstention", "The system admits it does not know instead of making a confident-looking guess."),
    ]
    t2 = doc.add_table(rows=1, cols=2)
    t2.style = "Table Grid"
    for c, text in zip(t2.rows[0].cells, ["Concept", "Meaning"], strict=True):
        set_cell_text(c, text, True, "FFFFFF")
        shade(c, "0F766E")
    for concept, meaning in concepts:
        cells = t2.add_row().cells
        set_cell_text(cells[0], concept, True)
        set_cell_text(cells[1], meaning)

    doc.add_heading("7. Your role as the learner", level=1)
    doc.add_paragraph("You are not expected to know everything before starting. You are expected to inspect the code, run commands, ask why, and attempt small exercises. I can implement with you, but I will explain the design and leave space for you to reason about the next step.")
    doc.add_paragraph("A successful session is not the session with the most code. It is the session where you can explain what changed, why it changed, and how we know it works.")

    doc.add_heading("8. First exercise", level=1)
    doc.add_paragraph("From the repository directory, run:")
    code = doc.add_paragraph()
    code.style = "No Spacing"
    run = code.add_run("uv run pytest\nuv run ruff check .\nuv run mypy src")
    run.font.name = "Menlo"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string("0F172A")
    doc.add_paragraph("Then answer these questions in your own words:")
    for q in [
        "What is the difference between detection and identification?",
        "Why are we starting with only 3–5 SKUs?",
        "What does unknown mean in a vision system?",
        "Why do we need measurements instead of only checking whether the model runs?",
        "Why does the dashboard come after the API and vision baseline?",
    ]:
        doc.add_paragraph(q, style="List Bullet")

    doc.add_heading("9. Next lesson", level=1)
    doc.add_paragraph("We will build a health endpoint and settings object. It will be the first real feature and will teach how the application starts, how an HTTP request reaches code, and how we test an external behavior.")
    doc.add_paragraph("Expected endpoint:")
    code = doc.add_paragraph()
    run = code.add_run("GET /healthz  →  {\"status\": \"ok\", \"version\": \"0.1.0\"}")
    run.font.name = "Menlo"
    run.font.size = Pt(9)

    doc.add_heading("10. Useful project files", level=1)
    for path, meaning in [
        ("README.md", "Project overview and commands"),
        ("docs/PLAN.md", "Technical architecture and roadmap"),
        ("src/shelfsense/", "Application code"),
        ("tests/", "Automated tests"),
        ("lessons/", "Short project lessons"),
        ("reference/", "Quick-reference material"),
        ("learning-records/", "What has been learned"),
    ]:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(path).bold = True
        p.add_run(" — " + meaning)

    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("ShelfSense · Learning by building · Local-first computer vision").font.size = Pt(8)

    doc.save(OUT / "ShelfSense-Learning-Plan.docx")


if __name__ == "__main__":
    pipeline_png()
    roadmap_png()
    make_doc()
    print(OUT / "ShelfSense-Learning-Plan.docx")
